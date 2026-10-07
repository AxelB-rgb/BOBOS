"""Topology from MSI, exact BIM nomenclature joins to CDE, and electrical readings."""
from __future__ import annotations
import json
import re
from collections import defaultdict
from pathlib import Path
import openpyxl
from .db import SCHEMA_PATH

MSI_NAME = '20220825 (On Dijon) Coordination MSI.xlsx'
CDE_NAME = 'export-cde-16_10_2024.xlsx'
FEED_RE = re.compile(r'^(Depart_\d+|TGBT_\d+)$')
SENSOR_RE = re.compile(r'^/(?:[^/]+)/((?:Depart|TGBT)_\d+)_Energie_Active$')
FLOORS = {'N0':'RDC','N1':'Etage 1','N2':'Etage 2','N3':'Etage 3','N4':'Etage 4','N5':'Terrasse','S1':'Sous-sol'}


def tokens(value):
    return sorted({v.strip() for v in str(value or '').split(';') if v.strip() and v.strip() != '???'})


def key(value):
    return str(value or '').strip().casefold()


def circuit_usage(tags):
    m = re.search(r"circuitUsage='([^']+)'", str(tags or ''))
    label = m.group(1) if m else ''
    text = label.casefold()
    if 'eclairage' in text or 'éclairage' in text:
        return 'Lighting'
    if any(v in text for v in ['ventilo','ventilation','cvc','climat','cta']):
        return 'Ventilation and Auxilaries'
    if 'chaud' in text:
        return 'Heating'
    if 'froid' in text:
        return 'Cooling'
    if 'distribution' in text:
        return 'Distribution'
    if 'prise' in text:
        return 'Sockets'
    return label or 'Other'


def read_reference(data_dir: Path):
    """No fuzzy join: preserve all duplicate CDE candidates and source row numbers."""
    cde = defaultdict(list)
    cde_path = data_dir / CDE_NAME
    if cde_path.exists():
        w = openpyxl.load_workbook(cde_path, read_only=True, data_only=True)
        try:
            s = w['TwinOps Referential']
            headers = next(s.iter_rows(max_row=1, values_only=True))
            for number, values in enumerate(s.iter_rows(min_row=2, values_only=True), 2):
                r = dict(zip(headers, values))
                if r['Type'] != 'Equipment' or not key(r.get('nomenclature')):
                    continue
                cde[key(r['nomenclature'])].append({
                    'id':r['TwinOps ID'], 'name':r['Name'], 'technical_path':r['Tech Path'],
                    'localisation':r['Localisation Path'], 'nomenclature':r['nomenclature'],
                    'source_row':number})
        finally:
            w.close()
    w = openpyxl.load_workbook(data_dir / MSI_NAME, read_only=True, data_only=True)
    try:
        structure = [(n, r) for n,r in enumerate(w['Structure'].iter_rows(min_row=5, values_only=True),5) if r[7]]
        equipment = {str(r[1]).strip():(n,r) for n,r in enumerate(w['Equipements site'].iter_rows(min_row=5, values_only=True),5) if r[1]}
        devices = {str(r[1]).strip():(n,r) for n,r in enumerate(w['Device site'].iter_rows(min_row=5, values_only=True),5) if r[1]}
        hierarchy = set()
        for r in w['Sous-Comptage'].iter_rows(min_row=5, values_only=True):
            if r[2] == 'Energie_Active' and r[4] == 'Energie_Active':
                for parent in tokens(r[1]):
                    for child in tokens(r[3]):
                        hierarchy.add((parent,child))
    finally:
        w.close()
    return structure, equipment, devices, cde, hierarchy


def load_reference(conn, data_dir):
    conn.executescript(SCHEMA_PATH.read_text(encoding='utf-8'))
    for table in ['electrical_import','electrical_hourly','electrical_meters','circuit_assets','circuit_rooms','meter_hierarchy','electrical_circuits']:
        conn.execute(f'DELETE FROM {table}')
    structure, eq, devices, cde, hierarchy = read_reference(data_dir)
    circuits = {name for name in eq if FEED_RE.fullmatch(name)}
    # Supply chains can pass through a controller/terminal. Stop at the first meter.
    supplies = {name:tokens(r[11]) for name,(_,r) in eq.items()}
    supplies.update({name:tokens(r[14]) for name,(_,r) in devices.items()})
    def feeds(name, visited=frozenset()):
        if name in circuits:
            return {name}
        if name in visited:
            return set()
        return {f for parent in supplies.get(name,[]) for f in feeds(parent,visited|{name})}
    services = defaultdict(set)
    links = defaultdict(set)
    locations = {}
    for number,r in structure:
        room = str(r[7]).strip()
        locations[room] = (FLOORS.get(str(r[6]),str(r[5] or '')),r[15],r[14])
        for col in [17,18,20,21,22,23]:
            for asset in tokens(r[col]):
                services[asset].add(room)
        for asset in tokens(r[19]):
            for feed in feeds(asset):
                links[(feed,room)].add(f'Structure!T{number}: {asset}')
        conn.execute('''INSERT INTO rooms(code,name,floor,zone,source) VALUES (?,?,?,?,?)
            ON CONFLICT(code) DO UPDATE SET floor=COALESCE(NULLIF(rooms.floor,''),excluded.floor),
            zone=COALESCE(NULLIF(rooms.zone,''),excluded.zone)''',
            (room,str(r[14] or room),locations[room][0],r[15],'msi'))
    for name in sorted(circuits):
        number,r = eq[name]
        conn.execute('INSERT INTO electrical_circuits VALUES (?,?,?,?,?,?)',
            (name,str(r[3] or name).strip(),circuit_usage(r[8]),r[7],r[6],
             json.dumps({'source':MSI_NAME,'sheet':'Equipements site','row':number,'tags':r[8]},ensure_ascii=False)))
    matched = ambiguous = conflicts = assets_count = 0
    for kind, records, room_col, bim_col, name_col in [('equipment',eq,6,2,3),('device',devices,7,2,3)]:
        for name,(number,r) in records.items():
            if name in circuits:
                continue
            associated = {f for parent in supplies.get(name,[]) for f in feeds(parent)}
            if not associated:
                continue
            matches = cde.get(key(r[bim_col]),[])
            matched += bool(matches)
            ambiguous += len(matches)>1
            hosted = str(r[room_col] or '').strip()
            cde_rooms = {m['localisation'].rsplit('/',1)[-1] for m in matches if m['localisation']}
            conflict = bool(hosted and cde_rooms and hosted not in cde_rooms)
            conflicts += conflict
            # Service relationships supersede equipment hosting (e.g. roof extractor).
            served = services.get(name,set())
            if not served:
                candidate = hosted or (next(iter(cde_rooms)) if len(cde_rooms)==1 else '')
                served = {candidate} if candidate in locations else set()
            sheet = 'Equipements site' if kind=='equipment' else 'Device site'
            payload = {'id':name,'kind':kind,'name':str(r[name_col] or name),
                       'bim_code':r[bim_col],'hosted_room':hosted or None,'served_rooms':sorted(served),
                       'cde_matches':matches,'cde_match_status':'ambiguous' if len(matches)>1 else 'exact' if matches else 'unmatched',
                       'localisation_conflict':conflict,'source':MSI_NAME,'sheet':sheet,'source_row':number}
            for feed in associated:
                conn.execute('INSERT OR REPLACE INTO circuit_assets VALUES (?,?,?)',(feed,name,json.dumps(payload,ensure_ascii=False)))
                for room in served:
                    links[(feed,room)].add(f'{sheet}!{number}: {name} → elecFeeds')
            assets_count += 1
    for (feed,room),provenance in links.items():
        conn.execute('INSERT INTO circuit_rooms VALUES (?,?,?)',(feed,room,json.dumps(sorted(provenance),ensure_ascii=False)))
    conn.executemany('INSERT INTO meter_hierarchy VALUES (?,?)',sorted(hierarchy))
    report = {'msi_rooms':len(structure),'rooms_with_elecFeeds':sum(bool(r[19]) for _,r in structure),
              'circuits':len(circuits),'circuits_with_rooms':len({f for f,_ in links}),
              'linked_assets':assets_count,'cde_matched_assets':matched,'cde_ambiguous_assets':ambiguous,
              'localisation_conflicts':conflicts,'sources':[MSI_NAME,CDE_NAME]}
    conn.execute('INSERT INTO electrical_import VALUES (?,?)',('reference',json.dumps(report,ensure_ascii=False)))
    conn.commit()
    return report


def has_electrical(conn):
    exists = conn.execute("SELECT 1 FROM sqlite_master WHERE name='electrical_import'").fetchone()
    return bool(exists and conn.execute("SELECT 1 FROM electrical_import WHERE key='complete'").fetchone())


def measured_parents(meters, hierarchy):
    """Exclude any measured ancestor of a measured descendant, including indirect ones."""
    children = defaultdict(set)
    for parent,child in hierarchy:
        children[parent].add(child)
    def descendants(parent,visited):
        result = set()
        for child in children[parent]-visited:
            result.add(child)
            result.update(descendants(child,visited|{child}))
        return result
    return {m for m in meters if (descendants(m,{m}) & (meters-{m}))}


def load_meter_energy(conn, data_dir):
    import pandas as pd
    report = load_reference(conn,data_dir)
    known = {r[0] for r in conn.execute('SELECT id FROM electrical_circuits')}
    usages = {r[0]:r[1] for r in conn.execute('SELECT id,usage FROM electrical_circuits')}
    last = {}
    discarded = set()
    seen = {}
    for path in sorted((data_dir/'energy_csv').glob('energy_*.csv')):
        parts = []
        for df in pd.read_csv(path,sep=';',usecols=['timestamp','display name','sensor id','value','resource','re2020Usage'],chunksize=250000):
            df = df[df['resource'].eq('Electricity')].copy()
            df['circuit_id'] = df['sensor id'].str.extract(SENSOR_RE,expand=False)
            discarded.update(df.loc[df['circuit_id'].isna(),'sensor id'].unique())
            df = df[df['circuit_id'].notna()]
            if not df.empty:
                parts.append(df)
        if not parts:
            continue
        df = pd.concat(parts,ignore_index=True)
        df['timestamp'] = pd.to_datetime(df['timestamp'],errors='coerce',format='mixed')
        df['value'] = pd.to_numeric(df['value'],errors='coerce')
        df = df.dropna(subset=['timestamp','value']).sort_values(['sensor id','timestamp'])
        df = df.drop_duplicates(['sensor id','timestamp'])
        df['delta'] = df.groupby('sensor id')['value'].diff()
        for idx,row in df.groupby('sensor id',sort=False).head(1).iterrows():
            if row['sensor id'] in last:
                df.at[idx,'delta'] = row['value']-last[row['sensor id']]
        df['reset'] = df['delta'].lt(0).astype(int)
        df['delta'] = df['delta'].clip(lower=0).fillna(0)
        last.update(df.groupby('sensor id')['value'].last().to_dict())
        df['hour'] = df['timestamp'].dt.floor('h').dt.strftime('%Y-%m-%d %H:%M:%S')
        hourly = df.groupby(['sensor id','hour'],as_index=False).agg(energy=('delta','sum'),samples=('value','size'),resets=('reset','sum'))
        conn.executemany('''INSERT INTO electrical_hourly VALUES (?,?,?,?,?)
            ON CONFLICT(sensor_id,hour) DO UPDATE SET energy_kwh=electrical_hourly.energy_kwh+excluded.energy_kwh,
            n_samples=electrical_hourly.n_samples+excluded.n_samples,resets=electrical_hourly.resets+excluded.resets''',
            list(hourly.itertuples(index=False,name=None)))
        for _,r in df.drop_duplicates('sensor id').iterrows():
            feed = r['circuit_id']
            if feed not in known:
                conn.execute('INSERT OR IGNORE INTO electrical_circuits VALUES (?,?,?,?,?,?)',
                    (feed,str(r['display name']),str(r['re2020Usage'] or 'Other'),None,None,json.dumps({'source':'CSV','mapping':'unmatched'})))
                known.add(feed)
            seen[r['sensor id']] = (feed,str(r['display name']),usages.get(feed,str(r['re2020Usage'] or 'Other')))
        conn.commit()
        print(f'  Électricité {path.name}: {len(hourly):,} heures-compteurs',flush=True)
    excluded = measured_parents({v[0] for v in seen.values()},[(r[0],r[1]) for r in conn.execute('SELECT * FROM meter_hierarchy')])
    # Multiple distinct sensors claiming the same circuit are not silently summed.
    by_feed = defaultdict(list)
    for sensor,v in seen.items():
        by_feed[v[0]].append(sensor)
    duplicates = {feed:ids for feed,ids in by_feed.items() if len(ids)>1}
    for sensor,(feed,label,usage) in seen.items():
        conn.execute('INSERT INTO electrical_meters VALUES (?,?,?,?,?)',(sensor,feed,label,usage,int(feed not in excluded and feed not in duplicates)))
    report.update({'measured_meters':len(seen),'parent_meters_excluded':len(excluded),
                   'duplicate_circuit_sensors':duplicates,'ignored_electricity_sensors':sorted(discarded),
                   'included_meters':sum(feed not in excluded and feed not in duplicates for feed,_,_ in seen.values())})
    conn.execute('INSERT OR REPLACE INTO electrical_import VALUES (?,?)',('reference',json.dumps(report,ensure_ascii=False)))
    conn.execute('INSERT OR REPLACE INTO electrical_import VALUES (?,?)',('complete','1'))
    conn.commit()
    return report


def circuits_report(conn,start,end,include_assets=True):
    if not has_electrical(conn):
        return {'available':False,'notice':'Réimporter les données pour activer le référentiel électrique.','circuits':[]}
    report = json.loads(conn.execute("SELECT value FROM electrical_import WHERE key='reference'").fetchone()[0])
    links = defaultdict(list)
    for r in conn.execute('SELECT cr.*,r.name,r.floor,r.zone,EXISTS(SELECT 1 FROM occupancy_hourly o WHERE o.room=cr.room) has_presence FROM circuit_rooms cr LEFT JOIN rooms r ON r.code=cr.room ORDER BY cr.room'):
        links[r['circuit_id']].append({'code':r['room'],'name':r['name'],'floor':r['floor'],'zone':r['zone'],'has_presence':bool(r['has_presence']),'provenance':json.loads(r['provenance'])})
    assets = defaultdict(list)
    if include_assets:
        for r in conn.execute('SELECT * FROM circuit_assets ORDER BY asset_id'):
            assets[r['circuit_id']].append(json.loads(r['payload']))
    rows = [dict(r) for r in conn.execute('''
        SELECT m.sensor_id,m.circuit_id,m.label,m.usage,m.included,c.panel,c.hosted_room,c.metadata,
            SUM(e.energy_kwh) energy_kwh,SUM(e.resets) resets,COUNT(e.hour) observed_hours
        FROM electrical_meters m JOIN electrical_circuits c ON c.id=m.circuit_id
        LEFT JOIN electrical_hourly e ON e.sensor_id=m.sensor_id AND e.hour>=? AND e.hour<?
        GROUP BY m.sensor_id ORDER BY SUM(e.energy_kwh) DESC,m.circuit_id''',(start,end))]
    by_usage = defaultdict(float)
    total = shared = dedicated = unassigned = 0.
    for r in rows:
        r['rooms'] = links[r['circuit_id']]
        r['presence_rooms'] = sum(room['has_presence'] for room in r['rooms'])
        r['assets'] = assets[r['circuit_id']]
        r['source'] = json.loads(r.pop('metadata'))
        r['mapping'] = 'dedicated' if len(r['rooms'])==1 else 'shared' if r['rooms'] else 'unmapped'
        r['included'] = bool(r['included'])
        value = r['energy_kwh']
        r['energy_kwh'] = round(value,3) if value is not None else None
        if r['included'] and value is not None:
            total += value
            by_usage[r['usage']] += value
            if r['mapping']=='dedicated': dedicated+=value
            elif r['mapping']=='shared': shared+=value
            else: unassigned+=value
    report['fully_monitored_circuits'] = sum(bool(c['rooms']) and c['presence_rooms']==len(c['rooms']) and c['included'] for c in rows)
    return {'available':True,'period':{'from':start,'to':end},'coverage':report,'circuits':rows,
            'summary':{'energy_kwh':round(total,3),'dedicated_kwh':round(dedicated,3),'shared_kwh':round(shared,3),
                       'unassigned_kwh':round(unassigned,3),'by_usage':dict((k,round(v,3)) for k,v in by_usage.items())},
            'notice':'Électricité mesurée aux compteurs sans descendants mesurés. Départs partagés non répartis entre salles. Ce sous-comptage ne constitue pas le total facturé du bâtiment.'}
